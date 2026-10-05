import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_callout_border(cell, border_color="0D5C3A", border_sz="36"):
    # thick left border, none for other sides
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                          f'<w:top w:val="none"/>'
                          f'<w:left w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
                          f'<w:bottom w:val="none"/>'
                          f'<w:right w:val="none"/>'
                          f'</w:tcBorders>')
    tcPr.append(tcBorders)

def set_table_borders(table, border_color="CBD5E1"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
                            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
                            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
                            f'<w:insideV w:val="none"/>'
                            f'<w:left w:val="none"/>'
                            f'<w:right w:val="none"/>'
                            f'</w:tblBorders>')
        tblPr[0].append(borders)

def build_document():
    doc = Document()

    # Set Margins (2.54 cm / 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)  # A4
        section.page_height = Inches(11.69)

    # Color definitions
    NAVY = RGBColor(15, 41, 66)        # #0F2942
    EMERALD = RGBColor(13, 92, 58)     # #0D5C3A
    GOLD = RGBColor(180, 83, 9)        # #B45309
    DARK_TEXT = RGBColor(30, 41, 59)   # #1E293B
    MUTED_TEXT = RGBColor(100, 116, 139) # #64748B
    WHITE = RGBColor(255, 255, 255)

    # Set normal style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = DARK_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    # Helper function for adding styled headings
    def add_h1(text, space_before=12, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text, space_before=8, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = EMERALD
        return p

    def add_h3(text, space_before=6, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_callout(title, text_content, border_hex="0D5C3A", bg_hex="F0FDF4"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.27)
        set_cell_background(cell, bg_hex)
        set_callout_border(cell, border_color=border_hex, border_sz="28")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if title:
            r_title = p.add_run(f"{title}\n")
            r_title.font.bold = True
            r_title.font.size = Pt(10.5)
            r_title.font.color.rgb = NAVY
        r_body = p.add_run(text_content)
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = DARK_TEXT

    # =========================================================================
    # HALAMAN 1: SAMPUL DAN IDENTITAS PROGRAM
    # =========================================================================
    p_kicker = doc.add_paragraph()
    p_kicker.paragraph_format.space_before = Pt(40)
    p_kicker.paragraph_format.space_after = Pt(8)
    p_kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_kicker = p_kicker.add_run("DOKUMEN PENGENALAN PROGRAM & BAHAN DISKUSI STRATEGIS")
    r_kicker.font.size = Pt(11)
    r_kicker.font.bold = True
    r_kicker.font.color.rgb = GOLD

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("KucingMu.online")
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(24)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Platform Digital Identitas, Rekam Medis, dan Kesejahteraan Kucing Terpadu")
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = EMERALD

    # Decorative Box for Cover Summary
    cover_box = doc.add_table(rows=1, cols=1)
    cover_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = cover_box.cell(0, 0)
    c_cell.width = Inches(5.8)
    set_cell_background(c_cell, "F8FAFC")
    set_callout_border(c_cell, border_color="0F2942", border_sz="20")
    set_cell_margins(c_cell, top=200, bottom=200, left=240, right=240)
    
    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    cp.paragraph_format.line_spacing = 1.2
    c_run1 = cp.add_run("Fokus Audiensi:\n")
    c_run1.font.bold = True
    c_run1.font.size = Pt(10.5)
    c_run1.font.color.rgb = NAVY
    c_run2 = cp.add_run(
        "1. Pemahaman latar belakang masalah dan solusi sistemik KucingMu.\n"
        "2. Nilai kemanfaatan bagi pemilik hewan, dokter hewan, masyarakat, dan persyarikatan.\n"
        "3. Peluang eksplorasi sinergi program edukasi, kesehatan (One Health), dan kesejahteraan satwa."
    )
    c_run2.font.size = Pt(10)
    c_run2.font.color.rgb = DARK_TEXT

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(120)
    p_meta.paragraph_format.space_after = Pt(2)
    p_meta.paragraph_format.line_spacing = 1.2
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    rm1 = p_meta.add_run("Pihak yang Dituju:\n")
    rm1.font.bold = True
    rm1.font.size = Pt(10)
    rm1.font.color.rgb = MUTED_TEXT
    
    rm2 = p_meta.add_run("Pimpinan Organisasi / Majelis / Lembaga Persyarikatan Muhammadiyah\n\n")
    rm2.font.bold = True
    rm2.font.size = Pt(11)
    rm2.font.color.rgb = NAVY

    rm3 = p_meta.add_run("Tim Penggagas & Pengembang:\n")
    rm3.font.bold = True
    rm3.font.size = Pt(10)
    rm3.font.color.rgb = MUTED_TEXT

    rm4 = p_meta.add_run("Denis & Tim Inisiator Platform KucingMu\n")
    rm4.font.bold = True
    rm4.font.size = Pt(11)
    rm4.font.color.rgb = DARK_TEXT

    rm5 = p_meta.add_run("Dokumen Diskusi dan Audiensi Strategis | Versi 1.0 (2026)")
    rm5.font.size = Pt(9.5)
    rm5.font.color.rgb = MUTED_TEXT

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 2: RINGKASAN EKSEKUTIF (EXECUTIVE SUMMARY)
    # =========================================================================
    add_h1("1. Ringkasan Eksekutif")

    p = doc.add_paragraph(
        "KucingMu (diakses melalui platform web KucingMu.online) adalah inisiatif platform digital "
        "yang dirancang untuk menjawab tantangan tata kelola identitas hewan peliharaan, pencatatan rekam medis "
        "terpusat, serta peningkatan literasi kesejahteraan hewan di Indonesia. Dalam konteks sosial keagamaan, "
        "kucing merupakan satwa yang sangat dekat dengan kehidupan sehari-hari masyarakat Muslim, termasuk warga persyarikatan. "
        "Namun, saat ini pengelolaan kesehatan dan kepemilikan kucing masih bersifat konvensional, terfragmentasi, "
        "dan minim integrasi data."
    )

    add_h2("Permasalahan Utama yang Dijawab")
    p = doc.add_paragraph(
        "Terdapat tiga celah utama dalam ekosistem perawatan kucing saat ini: "
        "(1) tidak adanya identitas digital resmi yang valid dan mudah diverifikasi publik, "
        "(2) rekam medis hewan yang tersebar di berbagai klinik tanpa riwayat berkelanjutan saat berpindah faskes, dan "
        "(3) rendahnya kesadaran pemilik mengenai risiko penyakit zoonosis (penyakit yang menular dari hewan ke manusia seperti rabies dan toksoplasmosis)."
    )

    add_h2("Solusi dan Fitur Utama")
    p = doc.add_paragraph(
        "KucingMu menghadirkan ekosistem terpadu berbasis peran (multi-role) yang mencakup:\n"
        "• Kartu Tanda Anggota KucingMu (KTAKuMu): Identitas digital kucing dengan Nomor Induk Anggota (NIAKuMu), kode QR unik, dan data biometrik.\n"
        "• Rekam Medis Elektronik (RME): Pencatatan riwayat pemeriksaan fisik, vaksinasi, pengobatan, dan riwayat klinis langsung oleh dokter hewan berwenang.\n"
        "• Sistem Verifikasi Berjenjang: Verifikasi berkas dan catatan medis oleh Verifikator resmi sebelum KTAKuMu diterbitkan.\n"
        "• Verifikasi Publik Cepat: Siapa pun dapat memindai QR Code untuk memeriksa keaslian data kucing dan validitas status medisnya."
    )

    add_h2("Nilai Manfaat Sosial dan Sinergi dengan Muhammadiyah")
    p = doc.add_paragraph(
        "Bagi masyarakat dan pemilik kucing, platform ini memberikan kepastian riwayat kesehatan serta kemudahan perawatan. "
        "Bagi praktisi medis hewan, platform ini menyediakan standar pencatatan medis digital yang efisien. "
        "Bagi persyarikatan Muhammadiyah, program ini membuka peluang penguatan dakwah bil hal di bidang lingkungan hidup "
        "(Majelis Lingkungan Hidup), kesehatan masyarakat berbasis One Health (Majelis Kesehatan Masyarakat/PKU), "
        "dan penguatan aksi kepedulian sosial (Lazismu)."
    )

    add_callout(
        "Tujuan Audiensi",
        "Dokumen ini disusun sebagai bahan diskusi awal untuk mengenalkan arsitektur program, menggali masukan dari pimpinan, "
        "serta menjajaki potensi uji coba (pilot project) bersama ekosistem persyarikatan yang relevan tanpa mengikat komitmen formal "
        "sebelum tercapai kesepahaman bersama."
    )

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 3: LATAR BELAKANG, URGENSI, VISI, DAN MISI
    # =========================================================================
    add_h1("2. Latar Belakang dan Urgensi Program")

    p = doc.add_paragraph(
        "Pertumbuhan populasi kucing peliharaan maupun kucing jalanan (stray cats) di wilayah perkotaan "
        "dan pemukiman terus meningkat pesat dalam beberapa tahun terakhir. Di lingkungan warga persyarikatan, "
        "kucing menjadi hewan peliharaan paling populer karena memiliki kedudukan istimewa dalam tradisi Islam "
        "sebagai satwa kesayangan yang suci dan patut disayangi (ihsan fi kulli syai')."
    )

    p = doc.add_paragraph(
        "Meskipun demikian, peningkatan adopsi kucing belum diimbangi oleh infrastruktur pendukung yang memadai. "
        "Berikut adalah pemetaan kondisi faktual di lapangan:"
    )

    # Table for Verified Problems vs Hypotheses
    tbl_prob = doc.add_table(rows=4, cols=2)
    tbl_prob.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_prob)

    headers = ["Kondisi Terverifikasi (Faktual di Lapangan)", "Asumsi / Kebutuhan yang Perlu Divalidasi"]
    for i, h in enumerate(headers):
        cell = tbl_prob.cell(0, i)
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(9.5)

    prob_data = [
        ("Buku kesehatan fisik (kertas) mudah hilang, rusak, atau tertinggal saat hewan membutuhkan pertolongan medis darurat.",
         "Tingkat kesediaan pemilik hewan untuk rutin memperbarui data digital secara mandiri."),
        ("Tidak adanya standarisasi pencatatan riwayat medis antar klinik mandiri, menyulitkan dokter hewan baru mengetahui riwayat terapi sebelumnya.",
         "Format rekam medis yang paling ringkas dan praktis bagi dokter hewan dalam praktik lapangan yang padat."),
        ("Rendahnya pemahaman umum mengenai pentingnya vaksinasi rutin, sterilisasi populasi, dan pencegahan penyakit menular (zoonosis).",
         "Kesiapan komunitas lokal untuk ikut memfasilitasi program edukasi dan sterilisasi terkoordinasi.")
    ]

    for row_idx, (c1, c2) in enumerate(prob_data, start=1):
        cell1 = tbl_prob.cell(row_idx, 0)
        cell2 = tbl_prob.cell(row_idx, 1)
        set_cell_margins(cell1, top=100, bottom=100, left=140, right=140)
        set_cell_margins(cell2, top=100, bottom=100, left=140, right=140)
        if row_idx % 2 == 1:
            set_cell_background(cell1, "F8FAFC")
            set_cell_background(cell2, "F8FAFC")
        
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(f"• {c1}")
        r1.font.size = Pt(9)

        p2 = cell2.paragraphs[0]
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(f"• {c2}")
        r2.font.size = Pt(9)

    add_h2("Visi, Misi, dan Nilai Dasar KucingMu")
    
    p = doc.add_paragraph()
    p.add_run("Visi Utama:\n").font.bold = True
    p.add_run("Mewujudkan ekosistem tata kelola kesehatan dan kesejahteraan kucing yang terpadu, bertanggung jawab, "
              "dan berbasis teknologi untuk mendukung kesehatan masyarakat dan kelestarian lingkungan.")

    p_misi = doc.add_paragraph()
    p_misi.add_run("Misi Program:\n").font.bold = True
    p_misi.add_run(
        "1. Menyediakan sistem identitas digital kucing yang valid, aman, dan mudah diakses oleh pemilik dan otoritas kesehatan.\n"
        "2. Mengintegrasikan rekam medis hewan peliharaan secara elektronik guna meningkatkan mutu layanan kesehatan hewan.\n"
        "3. Membangun kesadaran masyarakat mengenai prinsip kesejahteraan satwa (animal welfare) dan pencegahan penyakit zoonosis.\n"
        "4. Mendorong kolaborasi lintas sektor antara akademisi, praktisi medis hewan, komunitas sosial, dan organisasi keagamaan."
    )

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 4: GAMBARAN APLIKASI DAN FITUR UTAMA (BAGIAN 1)
    # =========================================================================
    add_h1("3. Gambaran Platform dan Fitur Utama (1/2)")

    p = doc.add_paragraph(
        "KucingMu.online dibangun dengan prinsip antarmuka yang ramah pengguna (user-friendly), ringan, "
        "dan dapat diakses dengan cepat melalui browser komputer maupun ponsel pintar tanpa perlu instalasi aplikasi berat. "
        "Sistem dirancang untuk melayani kebutuhan berbagai pihak dalam satu alur kerja yang terstruktur."
    )

    add_h2("1. Identitas Digital & Kartu Tanda Anggota (KTAKuMu)")
    p = doc.add_paragraph(
        "Setiap kucing yang didaftarkan memperoleh profil digital lengkap yang memuat informasi krusial:\n"
        "• Nomor Induk Anggota KucingMu (NIAKuMu): Nomor identifikasi unik berbasis wilayah pendaftaran dan urutan registrasi.\n"
        "• Profil Pemilik & Keterikatan Komunitas: Terhubung dengan data pemilik terverifikasi, termasuk opsi pencatatan Nomor Baku Muhammadiyah (NBM) bagi warga persyarikatan.\n"
        "• Parameter Fisik & Biometrik: Mencatat ras, warna bulu, jenis kelamin, pola tanda khas, dan jenis biometrik (seperti cetak hidung / nose print atau microchip jika terpasang).\n"
        "• Kartu Digital & Cetak: Format kartu resmi yang dilengkapi kode QR untuk verifikasi instan di klinik atau ruang publik."
    )

    add_h2("2. Rekam Medis Elektronik (RME) Terintegrasi")
    p = doc.add_paragraph(
        "Fitur ini memungkinkan dokter hewan mencatat hasil pemeriksaan klinis secara sistematis:\n"
        "• Pemeriksaan Fisik: Berat badan, suhu tubuh, kondisi umum, dan tanda vital kucing saat datang.\n"
        "• Riwayat Vaksinasi & Imunisasi: Pencatatan tanggal vaksin (Rabies, Tricat, Tetracat), masa berlaku, dan jadwal booster.\n"
        "• Diagnosis & Tindakan Medis: Catatan diagnosis dokter hewan, terapi yang diberikan, dan tindakan medis (misal sterilisasi/kastrasi).\n"
        "• Resep & Instruksi Perawatan: Petunjuk perawatan mandiri di rumah bagi pemilik hewan."
    )

    add_callout(
        "Pentingnya Kontinuitas Rekam Medis",
        "Saat pemilik kucing berpindah kota atau memeriksakan kucing ke klinik lain, dokter baru dapat membaca "
        "riwayat pengobatan sebelumnya secara transparan melalui otorisasi sistem. Hal ini mencegah kesalahan pemberian "
        "dosis obat ganda atau vaksinasi yang tidak perlu.",
        border_hex="0F2942",
        bg_hex="F1F5F9"
    )

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 5: GAMBARAN APLIKASI DAN FITUR UTAMA (BAGIAN 2)
    # =========================================================================
    add_h1("3. Gambaran Platform dan Fitur Utama (2/2)")

    add_h2("3. Arsitektur Multi-Peran (Multi-Role Ecosystem)")
    p = doc.add_paragraph(
        "Platform KucingMu menerapkan tata kelola akses berbasis peran untuk menjamin akuntabilitas data:"
    )

    # Table of Roles
    tbl_roles = doc.add_table(rows=5, cols=3)
    tbl_roles.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_roles)

    r_headers = ["Peran Pengguna", "Hak Akses & Tanggung Jawab Utama", "Status Kesiapan"]
    for i, h in enumerate(r_headers):
        cell = tbl_roles.cell(0, i)
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(9.5)

    roles_data = [
        ("Member (Pemilik Kucing)", "Mendaftarkan profil kucing, memantau riwayat rekam medis, mengunduh KTAKuMu digital, dan melihat jadwal perawatan.", "Sudah Berjalan (Live)"),
        ("Dokter Hewan (Veterinarian)", "Menginput hasil pemeriksaan fisik, rekam medis klinis, riwayat vaksin, resep, dan memberikan rekomendasi verifikasi.", "Sudah Berjalan (Live)"),
        ("Verifikator KTAKuMu", "Memeriksa kelengkapan berkas identitas dan hasil pemeriksaan medis sebelum kartu KTAKuMu resmi diterbitkan.", "Sudah Berjalan (Live)"),
        ("Admin & Pengelola Wilayah", "Mengelola master data wilayah, verifikasi anggota persyarikatan, memantau statistik populasi dan kesehatan.", "Sudah Berjalan (Live)")
    ]

    for row_idx, (c1, c2, c3) in enumerate(roles_data, start=1):
        cell1 = tbl_roles.cell(row_idx, 0)
        cell2 = tbl_roles.cell(row_idx, 1)
        cell3 = tbl_roles.cell(row_idx, 2)
        set_cell_margins(cell1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell2, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell3, top=80, bottom=80, left=120, right=120)
        if row_idx % 2 == 1:
            set_cell_background(cell1, "F8FAFC")
            set_cell_background(cell2, "F8FAFC")
            set_cell_background(cell3, "F8FAFC")

        cell1.paragraphs[0].add_run(c1).font.size = Pt(9)
        cell1.paragraphs[0].runs[0].font.bold = True
        cell2.paragraphs[0].add_run(c2).font.size = Pt(8.5)
        r3 = cell3.paragraphs[0].add_run(c3)
        r3.font.size = Pt(8.5)
        r3.font.bold = True
        r3.font.color.rgb = EMERALD

    add_h2("4. Matriks Status Kesiapan Fitur")
    p = doc.add_paragraph(
        "Untuk memberikan gambaran yang transparan kepada pimpinan, berikut adalah pengelompokan fitur berdasarkan status pengembangan:"
    )

    tbl_status = doc.add_table(rows=4, cols=2)
    tbl_status.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_status)

    s_headers = ["Kategori Tahapan", "Daftar Fitur dan Fungsionalitas"]
    for i, h in enumerate(s_headers):
        cell = tbl_status.cell(0, i)
        set_cell_background(cell, "0D5C3A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(9.5)

    status_data = [
        ("Fitur Sudah Berjalan (Live di Sistem)", 
         "Registrasi multi-role, pencatatan profil kucing lengkap, generator NIAKuMu & KTAKuMu otomatis, RME dokter hewan, sistem antrian verifikasi berjenjang, dan halaman verifikasi publik QR Code."),
        ("Fitur Dalam Tahap Finalisasi", 
         "Pengingat jadwal vaksinasi via notifikasi, filter pencarian rekam medis lanjutan, dan ringkasan riwayat medis siap cetak dalam format PDF."),
        ("Fitur Rencana Masa Depan (Roadmap)", 
         "Integrasi telekonsultasi dokter hewan online, peta sebaran klinik dan dokter mitra terdekat, serta modul donasi sterilisasi satwa jalanan terpadu.")
    ]

    for row_idx, (c1, c2) in enumerate(status_data, start=1):
        cell1 = tbl_status.cell(row_idx, 0)
        cell2 = tbl_status.cell(row_idx, 1)
        set_cell_margins(cell1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell2, top=80, bottom=80, left=120, right=120)
        if row_idx % 2 == 1:
            set_cell_background(cell1, "F8FAFC")
            set_cell_background(cell2, "F8FAFC")
        
        r1 = cell1.paragraphs[0].add_run(c1)
        r1.font.size = Pt(9)
        r1.font.bold = True
        cell2.paragraphs[0].add_run(c2).font.size = Pt(8.5)

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 6: MANFAAT DAN DAMPAK SOSIAL
    # =========================================================================
    add_h1("4. Manfaat dan Dampak Sosial yang Diharapkan")

    p = doc.add_paragraph(
        "KucingMu tidak dirancang semata-mata sebagai platform pencatatan administratif, melainkan sebagai "
        "sarana transformasi sosial untuk meningkatkan kualitas hidup bersama antara manusia dan hewan peliharaan "
        "sesuai dengan prinsip kesehatan masyarakat dan nilai-nilai Islam rahmatan lil 'alamin."
    )

    add_h2("Manfaat bagi Berbagai Pihak")
    
    p = doc.add_paragraph()
    p.add_run("1. Bagi Pemilik Kucing dan Keluarga:\n").font.bold = True
    p.add_run(
        "• Kepastian catatan medis yang rapi dan dapat diakses kapan pun saat dibutuhkan.\n"
        "• Meningkatkan rasa tanggung jawab (responsible pet ownership) dalam merawat hewan ciptaan Allah.\n"
        "• Memperoleh akses informasi kesehatan yang valid dan terhindar dari mitos perawatan yang membahayakan."
    )

    p = doc.add_paragraph()
    p.add_run("2. Bagi Dokter Hewan dan Mitra Klinik:\n").font.bold = True
    p.add_run(
        "• Standarisasi format pencatatan medis digital yang efisien dan menghemat waktu administrasi.\n"
        "• Kemudahan dalam melacak riwayat pasien berulang maupun pasien rujukan dari wilayah lain.\n"
        "• Ruang kontribusi profesional dalam mendukung program kesehatan hewan di lingkungan masyarakat."
    )

    p = doc.add_paragraph()
    p.add_run("3. Bagi Masyarakat Umum dan Lingkungan:\n").font.bold = True
    p.add_run(
        "• Menekan penyebaran penyakit menular zoonosis melalui pemantauan status vaksinasi yang terdata.\n"
        "• Mendorong pengendalian populasi kucing secara manusiawi melalui pendataan program sterilisasi.\n"
        "• Lingkungan pemukiman yang lebih bersih, sehat, dan bebas dari konflik sosial terkait satwa liar."
    )

    add_h2("Indikator Dampak yang Dapat Diukur (Key Performance Indicators)")
    
    tbl_kpi = doc.add_table(rows=5, cols=3)
    tbl_kpi.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_kpi)

    kpi_headers = ["Dimensi Dampak", "Indikator Kuantitatif Terukur", "Target Capaian Tahap Awal"]
    for i, h in enumerate(kpi_headers):
        cell = tbl_kpi.cell(0, i)
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(9.5)

    kpi_data = [
        ("Adopsi & Identitas Digital", "Jumlah kucing terdaftar dan memiliki kartu KTAKuMu resmi terverifikasi.", "500 - 1.000 kucing terdata pada wilayah percontohan."),
        ("Kesehatan Hewan", "Jumlah catatan rekam medis dan vaksinasi yang diinput oleh dokter hewan.", "100% kucing terdaftar memiliki rekam medis awal."),
        ("Jejaring Medis", "Jumlah dokter hewan dan klinik yang aktif menggunakan sistem.", "10 - 20 dokter hewan mitra pada tahap awal."),
        ("Edukasi & Literasi", "Jumlah kegiatan sosialisasi / edukasi perawatan kucing dan zoonosis.", "Minimal 2 sesi edukasi komunitas per semester.")
    ]

    for row_idx, (c1, c2, c3) in enumerate(kpi_data, start=1):
        cell1 = tbl_kpi.cell(row_idx, 0)
        cell2 = tbl_kpi.cell(row_idx, 1)
        cell3 = tbl_kpi.cell(row_idx, 2)
        set_cell_margins(cell1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell2, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell3, top=80, bottom=80, left=120, right=120)
        if row_idx % 2 == 1:
            set_cell_background(cell1, "F8FAFC")
            set_cell_background(cell2, "F8FAFC")
            set_cell_background(cell3, "F8FAFC")

        r1 = cell1.paragraphs[0].add_run(c1)
        r1.font.size = Pt(9)
        r1.font.bold = True
        cell2.paragraphs[0].add_run(c2).font.size = Pt(8.5)
        cell3.paragraphs[0].add_run(c3).font.size = Pt(8.5)

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 7: PELUANG SINERGI DENGAN MUHAMMADIYAH
    # =========================================================================
    add_h1("5. Peluang Sinergi dengan Ekosistem Muhammadiyah")

    p = doc.add_paragraph(
        "Sebagai gerakan Islam yang berakar kuat pada nilai tajdid (pembaruan) dan amal usaha di bidang kesehatan, "
        "sosial, dan pendidikan, Muhammadiyah memiliki infrastruktur organisasi yang sangat potensial "
        "untuk berkolaborasi dalam isu kesehatan terpadu (One Health) dan kesejahteraan lingkungan."
    )

    add_h2("Peta Eksplorasi Kolaborasi Antar Lembaga")

    tbl_sinergi = doc.add_table(rows=5, cols=2)
    tbl_sinergi.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sinergi)

    sin_headers = ["Unsur / Majelis / Lembaga", "Peluang Kolaborasi Program Strategis"]
    for i, h in enumerate(sin_headers):
        cell = tbl_sinergi.cell(0, i)
        set_cell_background(cell, "0D5C3A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(9.5)

    sinergi_data = [
        ("Majelis Lingkungan Hidup (MLH)", 
         "Kolaborasi kampanye kepedulian lingkungan dan kesejahteraan satwa urban, manajemen populasi kucing pemukiman ramah lingkungan, serta edukasi etika perlakuan terhadap satwa dalam perspektif Islam."),
        ("Majelis Kesehatan Masyarakat & RS PKU Muhammadiyah", 
         "Penerapan konsep One Health (keterkaitan kesehatan hewan, manusia, dan lingkungan) guna mengedukasi masyarakat terkait pencegahan penyakit zoonosis (rabies, toksoplasma, jamur kulit/ringworm)."),
        ("Perguruan Tinggi Muhammadiyah & 'Aisyiyah (PTMA)", 
         "Peluang kemitraan riset, praktikum mahasiswa, dan program pengabdian masyarakat melalui Fakultas Kedokteran Hewan atau Program Studi Peternakan/Biologi di lingkungan PTMA."),
        ("Lazismu (Lembaga Amil Zakat Muhammadiyah)", 
         "Eksplorasi program sosial tematik lingkungan, seperti fasilitasi pemeriksaan kesehatan gratis bagi kucing keluarga prasejahtera atau program sterilisasi kucing jalanan di area fasilitas publik.")
    ]

    for row_idx, (c1, c2) in enumerate(sinergi_data, start=1):
        cell1 = tbl_sinergi.cell(row_idx, 0)
        cell2 = tbl_sinergi.cell(row_idx, 1)
        set_cell_margins(cell1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell2, top=80, bottom=80, left=120, right=120)
        if row_idx % 2 == 1:
            set_cell_background(cell1, "F8FAFC")
            set_cell_background(cell2, "F8FAFC")
        
        r1 = cell1.paragraphs[0].add_run(c1)
        r1.font.size = Pt(9)
        r1.font.bold = True
        cell2.paragraphs[0].add_run(c2).font.size = Pt(8.5)

    add_callout(
        "Catatan Etika dan Independensi Dokumen",
        "Peluang sinergi di atas merupakan gagasan awal yang diajukan oleh tim pengembang untuk bahan pertimbangan "
        "dan diskusi bersama. KucingMu tidak menyatakan maupun mengklaim adanya kemitraan atau persetujuan formal "
        "sebelum tercapainya kesepakatan tertulis resmi dari pimpinan yang berwenang.",
        border_hex="B45309",
        bg_hex="FFFBEB"
    )

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 8: PROFIL PENGGAGAS DAN PENGEMBANG
    # =========================================================================
    add_h1("6. Profil Singkat Penggagas dan Pengembang")

    p = doc.add_paragraph(
        "Bagian ini disajikan sebagai informasi pendukung guna memberikan gambaran kepada pimpinan "
        "mengenai latar belakang keahlian, kapasitas teknis, dan komitmen tim yang menginisiasi platform KucingMu."
    )

    # Developer Profile Box
    prof_box = doc.add_table(rows=1, cols=1)
    prof_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_cell = prof_box.cell(0, 0)
    p_cell.width = Inches(6.27)
    set_cell_background(p_cell, "F8FAFC")
    set_callout_border(p_cell, border_color="0F2942", border_sz="24")
    set_cell_margins(p_cell, top=160, bottom=160, left=200, right=200)

    pp = p_cell.paragraphs[0]
    pp.paragraph_format.space_after = Pt(4)
    pp.paragraph_format.line_spacing = 1.15
    
    rp_name = pp.add_run("Denis\n")
    rp_name.font.bold = True
    rp_name.font.size = Pt(12)
    rp_name.font.color.rgb = NAVY

    rp_role = pp.add_run("Penggagas Inisiatif KucingMu & Praktisi Rekayasa Perangkat Lunak\n\n")
    rp_role.font.bold = True
    rp_role.font.size = Pt(9.5)
    rp_role.font.color.rgb = EMERALD

    rp_body = pp.add_run(
        "Latar Belakang & Pengalaman Profesional:\n"
        "• Berpengalaman dalam perancangan dan pengembangan sistem informasi digital, arsitektur basis data, "
        "serta aplikasi layanan publik berbasis web modern.\n"
        "• Memiliki rekam jejak aktif dalam pengelolaan infrastruktur digital sektor publik di lingkungan pemerintahan daerah "
        "(Diskominfotik Provinsi Lampung), dengan fokus pada keamanan data, keandalan sistem, dan kemudahan akses pengguna.\n"
        "• Memiliki minat mendalam terhadap pemanfaatan teknologi digital untuk solusi sosial kemasyarakatan, "
        "termasuk isu kesejahteraan satwa dan kesehatan publik."
    )
    rp_body.font.size = Pt(9)
    rp_body.font.color.rgb = DARK_TEXT

    add_h2("Kapasitas Teknis dan Komitmen Tata Kelola")
    p = doc.add_paragraph(
        "Dalam inisiatif KucingMu, tim pengembang berperan dalam:\n"
        "1. Pengembangan Perangkat Lunak: Membangun sistem yang aman, modular, dan terukur (scalable) menggunakan standar industri (Laravel, basis data relasional teroptimasi, enkripsi token).\n"
        "2. Pemeliharaan dan Keamanan: Menjamin ketersediaan platform, perlindungan privasi data pemilik hewan, dan backup sistem berkala.\n"
        "3. Keterbukaan Kolaborasi: Menyediakan integrasi teknis (API) dan siap menyesuaikan arsitektur sistem dengan kebutuhan regulasi maupun tata kelola organisasi mitra."
    )

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 9: ROADMAP DAN KEBERLANJUTAN PROGRAM
    # =========================================================================
    add_h1("7. Roadmap Pengembangan dan Keberlanjutan")

    p = doc.add_paragraph(
        "Rencana pengembangan platform disusun secara bertahap untuk memastikan setiap fase berjalan realistis, "
        "terukur, dan berlandaskan umpan balik pengguna di lapangan. Tahapan di bawah ini merupakan usulan kerja yang fleksibel."
    )

    tbl_road = doc.add_table(rows=4, cols=3)
    tbl_road.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_road)

    rd_headers = ["Fase Pengembangan", "Fokus Aktivitas Utama", "Kebutuhan & Luaran"]
    for i, h in enumerate(rd_headers):
        cell = tbl_road.cell(0, i)
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(9.5)

    road_data = [
        ("Fase 1 (Saat Ini): Fondasi & Stabilisasi", 
         "Pengujian fungsi sistem rekam medis, validasi alur penerbitan KTAKuMu, perapihan antarmuka responsif, dan penyusunan panduan pengguna.", 
         "Sistem inti siap pakai, dokumentasi operasional, dan materi sosialisasi."),
        ("Fase 2 (Usulan): Pilot Project Komunitas", 
         "Uji coba terbatas pada 1-2 wilayah/komunitas binaan, menjalin komunikasi dengan dokter hewan setempat, dan menghimpun masukan lapangan.", 
         "500 data kucing terverifikasi, evaluasi alur medis, dan laporan uji coba."),
        ("Fase 3 (Masa Depan): Ekspansi & Integrasi", 
         "Penyempurnaan modul telekonsultasi, perluasan jejaring klinik mitra, edukasi massal terstruktur, dan integrasi modul CSR/Lazismu.", 
         "Jejaring klinik mandiri, kegiatan edukasi rutin, dan dampak sosial terukur.")
    ]

    for row_idx, (c1, c2, c3) in enumerate(road_data, start=1):
        cell1 = tbl_road.cell(row_idx, 0)
        cell2 = tbl_road.cell(row_idx, 1)
        cell3 = tbl_road.cell(row_idx, 2)
        set_cell_margins(cell1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell2, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell3, top=80, bottom=80, left=120, right=120)
        if row_idx % 2 == 1:
            set_cell_background(cell1, "F8FAFC")
            set_cell_background(cell2, "F8FAFC")
            set_cell_background(cell3, "F8FAFC")

        r1 = cell1.paragraphs[0].add_run(c1)
        r1.font.size = Pt(9)
        r1.font.bold = True
        cell2.paragraphs[0].add_run(c2).font.size = Pt(8.5)
        cell3.paragraphs[0].add_run(c3).font.size = Pt(8.5)

    add_h2("Model Keberlanjutan Program (Sustainability)")
    p = doc.add_paragraph(
        "Untuk menjamin kelangsungan platform dalam jangka panjang, dirancang skema keberlanjutan non-eksploitatif:\n"
        "• Efisiensi Infrastruktur: Platform dibangun dengan arsitektur komputasi awan yang hemat sumber daya.\n"
        "• Nilai Tambah Layanan: Akses data dasar dan KTAKuMu digital tetap gratis untuk masyarakat luas, sementara biaya layanan fisik (seperti cetak kartu fisik PVC atau sertifikasi khusus) dapat menopang biaya server.\n"
        "• Gotong Royong Sosial: Peluang sinergi program sosial dengan lembaga filantropi untuk pendanaan kegiatan vaksinasi dan sterilisasi subsidi."
    )

    doc.add_page_break()

    # =========================================================================
    # HALAMAN 10: USULAN KOLABORASI DAN PENUTUP
    # =========================================================================
    add_h1("8. Usulan Bentuk Kolaborasi dan Penutup")

    p = doc.add_paragraph(
        "Sebagai penutup dari dokumen pengenalan ini, tim pengembang mengajukan beberapa poin konkret "
        "yang diharapkan dapat menjadi agenda pembahasan dalam sesi diskusi bersama pimpinan:"
    )

    add_h2("Poin-Poin Pokok Diskusi yang Diharapkan")
    
    p = doc.add_paragraph()
    p.add_run("1. Masukan & Arahan Strategis:\n").font.bold = True
    p.add_run(
        "Mendapatkan pandangan, nasihat, dan arahan pimpinan mengenai kesesuaian inisiatif KucingMu dengan agenda "
        "dakwah kemasyarakatan, lingkungan hidup, dan kesehatan di lingkungan persyarikatan."
    )

    p = doc.add_paragraph()
    p.add_run("2. Peluang Pilot Project Terbatas:\n").font.bold = True
    p.add_run(
        "Menjajaki kemungkinan pelaksanaan program percontohan pendaftaran dan pemeriksaan kesehatan kucing "
        "pada komunitas warga atau lingkungan ranting/cabang percontohan yang memiliki populasi kucing binaan."
    )

    p = doc.add_paragraph()
    p.add_run("3. Sinergi Edukasi & Sosialisasi Bersama:\n").font.bold = True
    p.add_run(
        "Menyelenggarakan webinar atau sosialisasi bersama mengenai kesehatan hewan keluarga, etika kepedulian satwa, "
        "dan pencegahan penyakit menular bagi warga masyarakat luas."
    )

    add_h2("Langkah Tindak Lanjut yang Diusulkan")

    tbl_followup = doc.add_table(rows=4, cols=2)
    tbl_followup.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_followup)

    fo_headers = ["Tahapan Tindak Lanjut", "Agenda / Aksi Konkret"]
    for i, h in enumerate(fo_headers):
        cell = tbl_followup.cell(0, i)
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        cp = cell.paragraphs[0]
        r = cp.add_run(h)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.size = Pt(9.5)

    fo_data = [
        ("Langkah 1: Sesi Diskusi & Demonstrasi", 
         "Pemaparan langsung sistem KucingMu.online kepada perwakilan majelis/lembaga terkait untuk demonstrasi alur kerja platform."),
        ("Langkah 2: Penyusunan Kerangka Kerja", 
         "Jika terdapat kesepahaman, menyusun kerangka kerja pelaksanaan uji coba dan pembagian peran teknis maupun edukatif."),
        ("Langkah 3: Evaluasi & Pengembangan Lanjutan", 
         "Melakukan evaluasi berkala terhadap hasil uji coba untuk merumuskan langkah kemitraan formal jangka panjang.")
    ]

    for row_idx, (c1, c2) in enumerate(fo_data, start=1):
        cell1 = tbl_followup.cell(row_idx, 0)
        cell2 = tbl_followup.cell(row_idx, 1)
        set_cell_margins(cell1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell2, top=80, bottom=80, left=120, right=120)
        if row_idx % 2 == 1:
            set_cell_background(cell1, "F8FAFC")
            set_cell_background(cell2, "F8FAFC")

        r1 = cell1.paragraphs[0].add_run(c1)
        r1.font.size = Pt(9)
        r1.font.bold = True
        cell2.paragraphs[0].add_run(c2).font.size = Pt(8.5)

    add_callout(
        "Penutup",
        "Besar harapan kami inisiatif KucingMu ini dapat memberikan manfaat nyata bagi kesehatan hewan peliharaan, "
        "ketenangan pemilik, dan kesehatan masyarakat luas, sekaligus menjadi bagian dari ikhtiar bersama dalam mewujudkan "
        "kebaikan bagi seluruh alam. Terima kasih atas perhatian dan kesempatan diskusi yang diberikan.",
        border_hex="0D5C3A",
        bg_hex="F0FDF4"
    )

    # Save to file
    output_filename = "Bahan Kucingmu.docx"
    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")

if __name__ == "__main__":
    build_document()
